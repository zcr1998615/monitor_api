"""
网络流量监控工具 - 核心模块

本模块提供使用 Playwright 从 Web 浏览器捕获和记录网络流量（API 调用）
的核心功能。

作者: 网络流量监控工具 v1.0
日期: 2026-04-23
"""

import json
import re
import pandas as pd
from playwright.sync_api import sync_playwright
import time
from datetime import datetime
import os
from urllib.parse import urlparse
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class TrafficMonitor:
    """
    流量监控类

    负责从浏览器实例拦截和记录网络请求及响应。使用 Playwright
    的浏览器自动化功能来捕获所有 fetch 和 xhr 请求。

    属性:
        target_url (str): 要监控的目标 URL
        output_file (str): 保存 Excel 输出文件的路径
        filter_keywords (list): 用于过滤请求的关键词列表
        exclude_domains (list): 排除的域名列表
        method_filter (list): 只捕获的 HTTP 方法列表
        status_filter (str): 状态码过滤 (fail/2xx/4xx/5xx/具体码)
        url_pattern (re.Pattern): URL 正则匹配模式
        append_mode (bool): 是否追加模式 (写入已有 Excel 的新 sheet)
        records (list): 用于存储捕获流量记录的列表
    """

    def __init__(self, target_url, output_file, filter_keywords=None,
                 exclude_domains=None, method_filter=None,
                 status_filter=None, url_pattern=None,
                 append_mode=False, generate_curl=False,
                 generate_docs=False):
        """
        初始化 TrafficMonitor 实例。

        参数:
            target_url (str): 要监控的目标 URL
            output_file (str): 保存 Excel 输出文件的路径
            filter_keywords (list, 可选): 关键词过滤列表
            exclude_domains (list, 可选): 排除的域名列表 (如 CDN)
            method_filter (list, 可选): 只捕获的 HTTP 方法列表 (如 ['GET', 'POST'])
            status_filter (str, 可选): 状态码过滤 (fail/2xx/4xx/5xx/具体码)
            url_pattern (str, 可选): URL 正则匹配
            append_mode (bool): 追加模式，写入已有 Excel 的新 sheet
            generate_curl (bool): 是否生成 cURL 命令文件
            generate_docs (bool): 是否生成接口文档
        """
        self.target_url = target_url
        self.output_file = output_file
        self.filter_keywords = filter_keywords or []
        self.exclude_domains = exclude_domains or []
        self.method_filter = [m.upper() for m in (method_filter or [])]
        self.status_filter = status_filter
        self.url_pattern = re.compile(url_pattern) if url_pattern else None
        self.append_mode = append_mode
        self._generate_curl = generate_curl
        self._generate_docs = generate_docs
        self.records = []
        self.stats = {"total": 0, "success": 0, "fail": 0, "total_duration": 0.0}

    def _attach_listeners(self, page):
        """为页面注册请求和响应事件监听器。"""
        page.on("request", self.handle_request)
        page.on("response", self.handle_response)

    def handle_request(self, request):
        """
        处理发出的 HTTP 请求。

        此回调在浏览器发出的每个网络请求时触发。它过滤请求以仅捕获 fetch 和 xhr
        类型（API 调用），并可按关键词、域名、HTTP 方法、正则进行过滤。

        参数:
            request: Playwright 的请求对象
        """
        if request.resource_type not in ["fetch", "xhr"]:
            return

        url = request.url

        if self.filter_keywords:
            if not any(kw in url for kw in self.filter_keywords):
                return

        if self.exclude_domains:
            try:
                domain = urlparse(url).netloc
                if any(d in domain for d in self.exclude_domains):
                    return
            except Exception:
                pass

        if self.method_filter:
            if request.method.upper() not in self.method_filter:
                return

        if self.url_pattern:
            if not self.url_pattern.search(url):
                return

        record = {
            "时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3],
            "方法": request.method,
            "URL": request.url,
            "请求头": json.dumps(request.headers, indent=2, ensure_ascii=False),
            "载荷": request.post_data or "",
            "状态": "待处理",
            "耗时(ms)": "",
            "响应头": "",
            "响应体": ""
        }
        request._traffic_record = record
        request._traffic_start_time = time.time()

    def handle_response(self, response):
        """
        处理收到的 HTTP 响应。

        此回调在浏览器收到的每个网络响应时触发。它将响应与相应的请求关联并更新记录。
        支持按状态码过滤。

        参数:
            response: Playwright 的响应对象
        """
        request = response.request
        if not hasattr(request, "_traffic_record"):
            return

        record = request._traffic_record
        record["状态"] = response.status
        duration = round((time.time() - request._traffic_start_time) * 1000, 2) if hasattr(request, "_traffic_start_time") else 0
        record["耗时(ms)"] = duration
        record["响应头"] = json.dumps(response.headers, indent=2, ensure_ascii=False)

        try:
            content_type = response.headers.get("content-type", "").lower()
            if "image" in content_type or "font" in content_type:
                record["响应体"] = "[二进制内容]"
            else:
                record["响应体"] = response.text()
        except Exception as e:
            record["响应体"] = f"[读取响应错误: {str(e)}]"

        if self.status_filter:
            status = response.status
            sf = self.status_filter.lower()
            skip = False
            if sf == "fail" and status < 400:
                skip = True
            elif sf == "4xx" and not (400 <= status < 500):
                skip = True
            elif sf == "5xx" and not (500 <= status < 600):
                skip = True
            elif sf == "2xx" and not (200 <= status < 300):
                skip = True
            elif sf.isdigit() and status != int(sf):
                skip = True
            if skip:
                print(f"[已过滤] [{record['方法']}] {record['URL']} -> {record['状态']}")
                return

        self.records.append(record)

        self.stats["total"] += 1
        if 200 <= response.status < 400:
            self.stats["success"] += 1
        else:
            self.stats["fail"] += 1
        self.stats["total_duration"] += duration
        avg = self.stats["total_duration"] / self.stats["total"]
        print(f"[{record['方法']}] {record['URL']} -> {record['状态']} ({duration}ms)")
        print(f"  统计: {self.stats['total']}个请求 | 成功:{self.stats['success']} 失败:{self.stats['fail']} | 平均耗时: {avg:.0f}ms")

    def start(self):
        """
        启动流量监控过程。

        启动无痕模式的 Chromium 浏览器，导航到目标 URL，并开始捕获网络流量。
        支持多标签页监控，所有新开标签页自动纳入监控。
        监控持续到用户关闭所有浏览器标签页。
        """
        with sync_playwright() as p:
            print("正在启动浏览器 (无痕模式)...")
            browser = p.chromium.launch(headless=False)
            context = browser.new_context()

            context.on("page", lambda new_page: self._attach_listeners(new_page))

            page = context.new_page()
            self._attach_listeners(page)

            print(f"正在打开网页: {self.target_url}")
            page.goto(self.target_url)

            print("\n" + "="*50)
            print("监控已就绪！(支持多标签页)")
            print("请在该浏览器窗口中进行您的手工操作。")
            print("完成录制后，只需【直接关闭浏览器窗口】即可。")
            print("="*50 + "\n")

            try:
                while True:
                    pages = context.pages
                    if not pages:
                        break
                    pages[0].wait_for_event('close', timeout=0)
            except Exception:
                pass

            print("\n浏览器已关闭，正在保存数据...")
            self.save_to_excel()
            self.save_to_har()

            if self._generate_curl:
                self.save_curl_commands()
            if self._generate_docs:
                self.save_api_docs()

            if self.stats["total"] > 0:
                print(f"\n本次监控统计:")
                print(f"   总请求数: {self.stats['total']}")
                print(f"   成功/失败: {self.stats['success']}/{self.stats['fail']}")
                print(f"   平均耗时: {self.stats['total_duration'] / self.stats['total']:.0f}ms")

    def save_to_excel(self):
        """将捕获的流量数据保存到 Excel 文件。支持追加模式（写入新 sheet）。"""
        if not self.records:
            print("未捕获到任何符合条件的接口请求。")
            return

        df = pd.DataFrame(self.records)
        df.sort_values(by="时间", inplace=True)

        output_dir = os.path.dirname(os.path.abspath(self.output_file))
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)

        if self.append_mode and os.path.exists(self.output_file):
            sheet_name = datetime.now().strftime("%H%M%S")
            with pd.ExcelWriter(self.output_file, engine='openpyxl', mode='a') as writer:
                df.to_excel(writer, sheet_name=sheet_name, index=False)
            self._beautify_excel(sheet_name=sheet_name)
            print(f"Excel 已追加 sheet [{sheet_name}] 至:")
        else:
            df.to_excel(self.output_file, index=False)
            self._beautify_excel()
            print(f"Excel 已保存至:")

        print(f"   {os.path.abspath(self.output_file)}")

    def _beautify_excel(self, sheet_name=None):
        """美化 Excel 表格：表头样式、自动列宽、状态码着色、URL 超链接。"""
        wb = load_workbook(self.output_file)
        ws = wb[sheet_name] if sheet_name else wb.active
        self._apply_styles(ws)
        wb.save(self.output_file)

    def _apply_styles(self, ws):
        ws.freeze_panes = "A2"

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_length = max(max_length, min(len(str(cell.value).split('\n')[0]), 60))
            ws.column_dimensions[col_letter].width = max(max_length + 2, 12)

        status_col = None
        for idx, cell in enumerate(ws[1], 1):
            if cell.value == "状态":
                status_col = idx
                break

        if status_col:
            green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            yellow = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            for row in ws.iter_rows(min_row=2):
                cell = row[status_col - 1]
                try:
                    status = int(cell.value)
                    if 200 <= status < 300:
                        cell.fill = green
                    elif 400 <= status < 500:
                        cell.fill = yellow
                    elif status >= 500:
                        cell.fill = red
                except (ValueError, TypeError):
                    pass

        url_col = None
        for idx, cell in enumerate(ws[1], 1):
            if cell.value == "URL":
                url_col = idx
                break

        if url_col:
            link_font = Font(color="0563C1", underline="single")
            for row in ws.iter_rows(min_row=2):
                cell = row[url_col - 1]
                if cell.value and str(cell.value).startswith("http"):
                    cell.hyperlink = str(cell.value)
                    cell.font = link_font

    def save_to_har(self):
        """将捕获的流量数据导出为标准 HAR (HTTP Archive) 格式。"""
        if not self.records:
            return

        har_file = os.path.splitext(self.output_file)[0] + ".har"
        entries = []

        for r in self.records:
            req_headers = []
            try:
                h = json.loads(r["请求头"])
                req_headers = [{"name": k, "value": v} for k, v in h.items()]
            except (json.JSONDecodeError, TypeError):
                pass

            resp_headers = []
            resp_content_type = "application/json"
            try:
                h = json.loads(r["响应头"])
                resp_headers = [{"name": k, "value": v} for k, v in h.items()]
                resp_content_type = h.get("content-type", "application/json").split(";")[0].strip()
            except (json.JSONDecodeError, TypeError):
                pass

            post_data = {}
            if r.get("载荷"):
                post_data = {"mimeType": "application/json", "text": str(r["载荷"])}

            duration = r.get("耗时(ms)", 0) or 0

            entry = {
                "startedDateTime": r["时间"].replace(" ", "T") + "+00:00",
                "time": duration,
                "request": {
                    "method": r["方法"],
                    "url": r["URL"],
                    "httpVersion": "HTTP/1.1",
                    "headers": req_headers,
                    "queryString": [],
                    "postData": post_data,
                    "headersSize": -1,
                    "bodySize": len(str(r.get("载荷", "")))
                },
                "response": {
                    "status": r["状态"] if isinstance(r["状态"], int) else 0,
                    "statusText": "",
                    "httpVersion": "HTTP/1.1",
                    "headers": resp_headers,
                    "content": {
                        "size": len(str(r.get("响应体", ""))),
                        "mimeType": resp_content_type,
                        "text": str(r.get("响应体", ""))
                    },
                    "headersSize": -1,
                    "bodySize": len(str(r.get("响应体", "")))
                },
                "cache": {},
                "timings": {
                    "send": 0,
                    "wait": duration,
                    "receive": 0
                }
            }
            entries.append(entry)

        har = {
            "log": {
                "version": "1.2",
                "creator": {"name": "Network Traffic Monitor", "version": "1.0"},
                "entries": entries
            }
        }

        with open(har_file, "w", encoding="utf-8") as f:
            json.dump(har, f, indent=2, ensure_ascii=False)

        print(f"HAR 已保存至:")
        print(f"   {os.path.abspath(har_file)}")

    def save_curl_commands(self):
        """将每条请求转换为 cURL 命令并保存到文件。"""
        if not self.records:
            return

        curl_file = os.path.splitext(self.output_file)[0] + "_curl.sh"
        lines = [
            "#!/bin/bash",
            f"# cURL 命令 - 生成于 {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"# 来源: {self.target_url}",
            f"# 共 {len(self.records)} 条请求",
            ""
        ]

        for i, r in enumerate(self.records, 1):
            parts = [f"curl -X {r['方法']}"]

            try:
                headers = json.loads(r["请求头"])
                for k, v in headers.items():
                    if k.startswith(":") or k.lower() in ("host", "connection", "content-length"):
                        continue
                    v_escaped = v.replace("'", "'\\''")
                    parts.append(f"  -H '{k}: {v_escaped}'")
            except (json.JSONDecodeError, TypeError):
                pass

            if r.get("载荷"):
                payload = str(r["载荷"]).replace("'", "'\\''")
                parts.append(f"  -d '{payload}'")

            url_escaped = r["URL"].replace("'", "'\\''")
            parts.append(f"  '{url_escaped}'")

            lines.append(f"# [{i}] {r['方法']} -> {r['状态']}")
            lines.append(" \\\n".join(parts))
            lines.append("")

        with open(curl_file, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))

        print(f"cURL 命令已保存至:")
        print(f"   {os.path.abspath(curl_file)}")

    def save_api_docs(self):
        """自动生成接口文档：按域名+路径去重，输出 Markdown 格式。"""
        if not self.records:
            return

        docs_file = os.path.splitext(self.output_file)[0] + "_api_docs.md"

        endpoints = {}
        for r in self.records:
            try:
                parsed = urlparse(r["URL"])
                key = (r["方法"], parsed.netloc, parsed.path)
                if key not in endpoints:
                    endpoints[key] = r
            except Exception:
                pass

        lines = [
            "# 接口文档",
            "",
            f"自动生成于: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"来源: {self.target_url}",
            f"接口总数: {len(endpoints)}",
            "",
            "---",
            ""
        ]

        by_domain = {}
        for (method, domain, path), record in endpoints.items():
            by_domain.setdefault(domain, []).append((method, path, record))

        for domain in sorted(by_domain.keys()):
            items = by_domain[domain]
            lines.append(f"## {domain}")
            lines.append("")

            for method, path, r in sorted(items, key=lambda x: x[1]):
                lines.append(f"### {method} {path}")
                lines.append("")
                lines.append(f"- **状态码**: {r['状态']}")
                lines.append(f"- **耗时**: {r.get('耗时(ms)', 'N/A')}ms")

                if r.get("载荷"):
                    lines.append("- **请求体**:")
                    try:
                        payload = json.loads(r["载荷"])
                        lines.append("  ```json")
                        lines.append(f"  {json.dumps(payload, indent=2, ensure_ascii=False)}")
                        lines.append("  ```")
                    except (json.JSONDecodeError, TypeError):
                        lines.append("  ```")
                        lines.append(f"  {r['载荷']}")
                        lines.append("  ```")

                body = r.get("响应体", "")
                if body and body not in ("[二进制内容]", ""):
                    lines.append("- **响应体**:")
                    try:
                        parsed_body = json.loads(body)
                        body_str = json.dumps(parsed_body, indent=2, ensure_ascii=False)
                        if len(body_str) > 500:
                            body_str = body_str[:500] + "\n  ... (已截断)"
                        lines.append("  ```json")
                        lines.append(f"  {body_str}")
                        lines.append("  ```")
                    except (json.JSONDecodeError, TypeError):
                        body_str = str(body)
                        if len(body_str) > 500:
                            body_str = body_str[:500] + "... (已截断)"
                        lines.append("  ```")
                        lines.append(f"  {body_str}")
                        lines.append("  ```")

                lines.append("")

        with open(docs_file, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))

        print(f"接口文档已保存至:")
        print(f"   {os.path.abspath(docs_file)}")


def compare_recordings(file1, file2, output=None):
    """
    对比两次录制文件，生成差异报告。

    参数:
        file1 (str): 第一个录制文件路径 (xlsx)
        file2 (str): 第二个录制文件路径 (xlsx)
        output (str, 可选): 输出报告路径，默认打印到控制台
    """
    df1 = pd.read_excel(file1)
    df2 = pd.read_excel(file2)

    is_zh = "方法" in df1.columns
    method_col = "方法" if is_zh else "Method"
    url_col = "URL"
    status_col = "状态" if is_zh else "Status"
    payload_col = "载荷" if is_zh else "Payload"
    body_col = "响应体" if is_zh else "Response Body"

    def extract_endpoints(df):
        endpoints = {}
        for _, row in df.iterrows():
            try:
                parsed = urlparse(str(row[url_col]))
                key = (str(row[method_col]), parsed.netloc, parsed.path)
                if key not in endpoints:
                    endpoints[key] = row
            except Exception:
                pass
        return endpoints

    ep1 = extract_endpoints(df1)
    ep2 = extract_endpoints(df2)

    keys1 = set(ep1.keys())
    keys2 = set(ep2.keys())

    added = keys2 - keys1
    removed = keys1 - keys2
    common = keys1 & keys2

    changed = []
    for key in sorted(common):
        r1, r2 = ep1[key], ep2[key]
        diffs = []
        if str(r1[status_col]) != str(r2[status_col]):
            diffs.append(f"状态码: {r1[status_col]} -> {r2[status_col]}")
        if str(r1.get(payload_col, "")) != str(r2.get(payload_col, "")):
            diffs.append("请求体发生变化")
        try:
            b1 = set(json.loads(str(r1[body_col])).keys())
            b2 = set(json.loads(str(r2[body_col])).keys())
            if b1 != b2:
                new_keys = b2 - b1
                del_keys = b1 - b2
                if new_keys:
                    diffs.append(f"响应新增字段: {', '.join(new_keys)}")
                if del_keys:
                    diffs.append(f"响应移除字段: {', '.join(del_keys)}")
        except Exception:
            pass
        if diffs:
            changed.append((key, diffs))

    lines = [
        "# 接口对比报告",
        "",
        f"文件A: {os.path.basename(file1)}",
        f"文件B: {os.path.basename(file2)}",
        f"对比时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
        "---",
        "",
        f"## 新增接口 ({len(added)}个，在B中有但A中没有)",
        ""
    ]

    if added:
        for method, domain, path in sorted(added):
            lines.append(f"- **{method}** `{domain}{path}`")
    else:
        lines.append("无")
    lines.append("")

    lines.append(f"## 移除接口 ({len(removed)}个，在A中有但B中没有)")
    lines.append("")
    if removed:
        for method, domain, path in sorted(removed):
            lines.append(f"- **{method}** `{domain}{path}`")
    else:
        lines.append("无")
    lines.append("")

    lines.append(f"## 变更接口 ({len(changed)}个)")
    lines.append("")
    if changed:
        for (method, domain, path), diffs in changed:
            lines.append(f"### {method} `{domain}{path}`")
            for d in diffs:
                lines.append(f"- {d}")
            lines.append("")
    else:
        lines.append("无")

    report = "\n".join(lines)

    if output:
        output_dir = os.path.dirname(os.path.abspath(output))
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir)
        with open(output, "w", encoding="utf-8") as f:
            f.write(report)
        print(f"对比报告已保存至: {os.path.abspath(output)}")
    else:
        print(report)

    return report
