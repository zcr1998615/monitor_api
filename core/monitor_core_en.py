"""
Network Traffic Monitor - Core Module

This module provides the core functionality for capturing and recording
network traffic (API calls) from web browsers using Playwright.

Author: Network Traffic Monitor v1.0
Date: 2026-04-23
"""

import json
import re
import pandas as pd
from playwright.sync_api import sync_playwright
import time
from datetime import datetime
import os
from urllib.parse import urlparse
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class TrafficMonitor:
    """
    Traffic Monitor Class

    Responsible for intercepting and recording network requests and responses
    from a browser instance. Uses Playwright's browser automation capabilities
    to capture all fetch and xhr requests.

    Attributes:
        target_url (str): The URL to monitor
        output_file (str): Path to save the Excel output file
        filter_keywords (list): Keywords to filter requests
        exclude_domains (list): Domains to exclude from capture
        method_filter (list): HTTP methods to capture
        status_filter (str): Status code filter (fail/2xx/4xx/5xx/specific code)
        url_pattern (re.Pattern): Regex pattern for URL matching
        append_mode (bool): Whether to append as new sheet to existing Excel
        records (list): List to store captured traffic records
    """

    def __init__(self, target_url, output_file, filter_keywords=None,
                 exclude_domains=None, method_filter=None,
                 status_filter=None, url_pattern=None,
                 append_mode=False, generate_curl=False,
                 generate_docs=False):
        """
        Initialize the TrafficMonitor instance.

        Args:
            target_url (str): The URL to monitor
            output_file (str): Path to save the Excel output file
            filter_keywords (list, optional): Keywords to filter requests
            exclude_domains (list, optional): Domains to exclude (e.g. CDN)
            method_filter (list, optional): HTTP methods to capture (e.g. ['GET', 'POST'])
            status_filter (str, optional): Status code filter (fail/2xx/4xx/5xx/specific code)
            url_pattern (str, optional): Regex pattern for URL matching
            append_mode (bool): Append as new sheet to existing Excel
            generate_curl (bool): Whether to generate cURL commands file
            generate_docs (bool): Whether to generate API documentation
        """
        self.target_url = target_url
        self.output_file = output_file
        self.filter_keywords = filter_keywords or []
        self.exclude_domains = exclude_domains or []
        self.method_filter = [m.upper() for m in (method_filter or [])]
        self.status_filter = status_filter
        self.url_pattern = re.compile(url_pattern) if url_pattern else None
        self.append_mode = append_mode
        self._generate_curl = generate_curl
        self._generate_docs = generate_docs
        self.records = []
        self.stats = {"total": 0, "success": 0, "fail": 0, "total_duration": 0.0}

    def _attach_listeners(self, page):
        """Attach request and response event listeners to a page."""
        page.on("request", self.handle_request)
        page.on("response", self.handle_response)

    def handle_request(self, request):
        """
        Handle outgoing HTTP requests.

        This callback is triggered for every network request made by the browser.
        It filters requests to only capture fetch and xhr types (API calls),
        and supports keyword, domain, method, and regex filtering.

        Args:
            request: Playwright's request object
        """
        if request.resource_type not in ["fetch", "xhr"]:
            return

        url = request.url

        if self.filter_keywords:
            if not any(kw in url for kw in self.filter_keywords):
                return

        if self.exclude_domains:
            try:
                domain = urlparse(url).netloc
                if any(d in domain for d in self.exclude_domains):
                    return
            except Exception:
                pass

        if self.method_filter:
            if request.method.upper() not in self.method_filter:
                return

        if self.url_pattern:
            if not self.url_pattern.search(url):
                return

        record = {
            "Time": datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3],
            "Method": request.method,
            "URL": request.url,
            "Request Headers": json.dumps(request.headers, indent=2, ensure_ascii=False),
            "Payload": request.post_data or "",
            "Status": "Pending",
            "Duration(ms)": "",
            "Response Headers": "",
            "Response Body": ""
        }
        request._traffic_record = record
        request._traffic_start_time = time.time()

    def handle_response(self, response):
        """
        Handle incoming HTTP responses.

        This callback is triggered for every network response received by the browser.
        It associates the response with the corresponding request and updates the record.
        Supports status code filtering.

        Args:
            response: Playwright's response object
        """
        request = response.request
        if not hasattr(request, "_traffic_record"):
            return

        record = request._traffic_record
        record["Status"] = response.status
        duration = round((time.time() - request._traffic_start_time) * 1000, 2) if hasattr(request, "_traffic_start_time") else 0
        record["Duration(ms)"] = duration
        record["Response Headers"] = json.dumps(response.headers, indent=2, ensure_ascii=False)

        try:
            content_type = response.headers.get("content-type", "").lower()
            if "image" in content_type or "font" in content_type:
                record["Response Body"] = "[Binary Content]"
            else:
                record["Response Body"] = response.text()
        except Exception as e:
            record["Response Body"] = f"[Error reading response: {str(e)}]"

        if self.status_filter:
            status = response.status
            sf = self.status_filter.lower()
            skip = False
            if sf == "fail" and status < 400:
                skip = True
            elif sf == "4xx" and not (400 <= status < 500):
                skip = True
            elif sf == "5xx" and not (500 <= status < 600):
                skip = True
            elif sf == "2xx" and not (200 <= status < 300):
                skip = True
            elif sf.isdigit() and status != int(sf):
                skip = True
            if skip:
                print(f"[filtered] [{record['Method']}] {record['URL']} -> {record['Status']}")
                return

        self.records.append(record)

        self.stats["total"] += 1
        if 200 <= response.status < 400:
            self.stats["success"] += 1
        else:
            self.stats["fail"] += 1
        self.stats["total_duration"] += duration
        avg = self.stats["total_duration"] / self.stats["total"]
        print(f"[{record['Method']}] {record['URL']} -> {record['Status']} ({duration}ms)")
        print(f"  Stats: {self.stats['total']} requests | OK:{self.stats['success']} FAIL:{self.stats['fail']} | Avg: {avg:.0f}ms")

    def start(self):
        """
        Start the traffic monitoring process.

        Launches a Chromium browser in incognito mode, navigates to the target URL,
        and begins capturing network traffic. Supports multi-tab monitoring.
        Monitoring continues until the user closes all browser tabs.
        """
        with sync_playwright() as p:
            print("Launching browser (incognito mode)...")
            browser = p.chromium.launch(headless=False)
            context = browser.new_context()

            context.on("page", lambda new_page: self._attach_listeners(new_page))

            page = context.new_page()
            self._attach_listeners(page)

            print(f"Opening webpage: {self.target_url}")
            page.goto(self.target_url)

            print("\n" + "="*50)
            print("Monitoring is ready! (multi-tab supported)")
            print("Please perform your operations in the browser window.")
            print("After completing, simply CLOSE the browser window.")
            print("="*50 + "\n")

            try:
                while True:
                    pages = context.pages
                    if not pages:
                        break
                    pages[0].wait_for_event('close', timeout=0)
            except Exception:
                pass

            print("\nBrowser closed, saving data...")
            self.save_to_excel()
            self.save_to_har()

            if self._generate_curl:
                self.save_curl_commands()
            if self._generate_docs:
                self.save_api_docs()

            if self.stats["total"] > 0:
                print(f"\nSession Stats:")
                print(f"   Total requests: {self.stats['total']}")
                print(f"   Success/Fail: {self.stats['success']}/{self.stats['fail']}")
                print(f"   Avg duration: {self.stats['total_duration'] / self.stats['total']:.0f}ms")

    def save_to_excel(self):
        """Save captured traffic data to Excel. Supports append mode (new sheet)."""
        if not self.records:
            print("No qualifying API requests were captured.")
            return

        df = pd.DataFrame(self.records)
        df.sort_values(by="Time", inplace=True)

        output_dir = os.path.dirname(os.path.abspath(self.output_file))
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)

        if self.append_mode and os.path.exists(self.output_file):
            sheet_name = datetime.now().strftime("%H%M%S")
            with pd.ExcelWriter(self.output_file, engine='openpyxl', mode='a') as writer:
                df.to_excel(writer, sheet_name=sheet_name, index=False)
            self._beautify_excel(sheet_name=sheet_name)
            print(f"Excel appended sheet [{sheet_name}] to:")
        else:
            df.to_excel(self.output_file, index=False)
            self._beautify_excel()
            print(f"Excel saved to:")

        print(f"   {os.path.abspath(self.output_file)}")

    def _beautify_excel(self, sheet_name=None):
        """Beautify Excel: header style, auto column width, status coloring, URL hyperlinks."""
        wb = load_workbook(self.output_file)
        ws = wb[sheet_name] if sheet_name else wb.active
        self._apply_styles(ws)
        wb.save(self.output_file)

    def _apply_styles(self, ws):
        ws.freeze_panes = "A2"

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_length = max(max_length, min(len(str(cell.value).split('\n')[0]), 60))
            ws.column_dimensions[col_letter].width = max(max_length + 2, 12)

        status_col = None
        for idx, cell in enumerate(ws[1], 1):
            if cell.value == "Status":
                status_col = idx
                break

        if status_col:
            green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            yellow = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            for row in ws.iter_rows(min_row=2):
                cell = row[status_col - 1]
                try:
                    status = int(cell.value)
                    if 200 <= status < 300:
                        cell.fill = green
                    elif 400 <= status < 500:
                        cell.fill = yellow
                    elif status >= 500:
                        cell.fill = red
                except (ValueError, TypeError):
                    pass

        url_col = None
        for idx, cell in enumerate(ws[1], 1):
            if cell.value == "URL":
                url_col = idx
                break

        if url_col:
            link_font = Font(color="0563C1", underline="single")
            for row in ws.iter_rows(min_row=2):
                cell = row[url_col - 1]
                if cell.value and str(cell.value).startswith("http"):
                    cell.hyperlink = str(cell.value)
                    cell.font = link_font

    def save_to_har(self):
        """Export captured traffic data to standard HAR (HTTP Archive) format."""
        if not self.records:
            return

        har_file = os.path.splitext(self.output_file)[0] + ".har"
        entries = []

        for r in self.records:
            req_headers = []
            try:
                h = json.loads(r["Request Headers"])
                req_headers = [{"name": k, "value": v} for k, v in h.items()]
            except (json.JSONDecodeError, TypeError):
                pass

            resp_headers = []
            resp_content_type = "application/json"
            try:
                h = json.loads(r["Response Headers"])
                resp_headers = [{"name": k, "value": v} for k, v in h.items()]
                resp_content_type = h.get("content-type", "application/json").split(";")[0].strip()
            except (json.JSONDecodeError, TypeError):
                pass

            post_data = {}
            if r.get("Payload"):
                post_data = {"mimeType": "application/json", "text": str(r["Payload"])}

            duration = r.get("Duration(ms)", 0) or 0

            entry = {
                "startedDateTime": r["Time"].replace(" ", "T") + "+00:00",
                "time": duration,
                "request": {
                    "method": r["Method"],
                    "url": r["URL"],
                    "httpVersion": "HTTP/1.1",
                    "headers": req_headers,
                    "queryString": [],
                    "postData": post_data,
                    "headersSize": -1,
                    "bodySize": len(str(r.get("Payload", "")))
                },
                "response": {
                    "status": r["Status"] if isinstance(r["Status"], int) else 0,
                    "statusText": "",
                    "httpVersion": "HTTP/1.1",
                    "headers": resp_headers,
                    "content": {
                        "size": len(str(r.get("Response Body", ""))),
                        "mimeType": resp_content_type,
                        "text": str(r.get("Response Body", ""))
                    },
                    "headersSize": -1,
                    "bodySize": len(str(r.get("Response Body", "")))
                },
                "cache": {},
                "timings": {
                    "send": 0,
                    "wait": duration,
                    "receive": 0
                }
            }
            entries.append(entry)

        har = {
            "log": {
                "version": "1.2",
                "creator": {"name": "Network Traffic Monitor", "version": "1.0"},
                "entries": entries
            }
        }

        with open(har_file, "w", encoding="utf-8") as f:
            json.dump(har, f, indent=2, ensure_ascii=False)

        print(f"HAR saved to:")
        print(f"   {os.path.abspath(har_file)}")

    def save_curl_commands(self):
        """Convert each request to a cURL command and save to file."""
        if not self.records:
            return

        curl_file = os.path.splitext(self.output_file)[0] + "_curl.sh"
        lines = [
            "#!/bin/bash",
            f"# cURL commands - generated {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"# Source: {self.target_url}",
            f"# Total: {len(self.records)} requests",
            ""
        ]

        for i, r in enumerate(self.records, 1):
            parts = [f"curl -X {r['Method']}"]

            try:
                headers = json.loads(r["Request Headers"])
                for k, v in headers.items():
                    if k.startswith(":") or k.lower() in ("host", "connection", "content-length"):
                        continue
                    v_escaped = v.replace("'", "'\\''")
                    parts.append(f"  -H '{k}: {v_escaped}'")
            except (json.JSONDecodeError, TypeError):
                pass

            if r.get("Payload"):
                payload = str(r["Payload"]).replace("'", "'\\''")
                parts.append(f"  -d '{payload}'")

            url_escaped = r["URL"].replace("'", "'\\''")
            parts.append(f"  '{url_escaped}'")

            lines.append(f"# [{i}] {r['Method']} -> {r['Status']}")
            lines.append(" \\\n".join(parts))
            lines.append("")

        with open(curl_file, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))

        print(f"cURL commands saved to:")
        print(f"   {os.path.abspath(curl_file)}")

    def save_api_docs(self):
        """Auto-generate API documentation: deduplicate by domain+path, output Markdown."""
        if not self.records:
            return

        docs_file = os.path.splitext(self.output_file)[0] + "_api_docs.md"

        endpoints = {}
        for r in self.records:
            try:
                parsed = urlparse(r["URL"])
                key = (r["Method"], parsed.netloc, parsed.path)
                if key not in endpoints:
                    endpoints[key] = r
            except Exception:
                pass

        lines = [
            "# API Documentation",
            "",
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"Source: {self.target_url}",
            f"Total endpoints: {len(endpoints)}",
            "",
            "---",
            ""
        ]

        by_domain = {}
        for (method, domain, path), record in endpoints.items():
            by_domain.setdefault(domain, []).append((method, path, record))

        for domain in sorted(by_domain.keys()):
            items = by_domain[domain]
            lines.append(f"## {domain}")
            lines.append("")

            for method, path, r in sorted(items, key=lambda x: x[1]):
                lines.append(f"### {method} {path}")
                lines.append("")
                lines.append(f"- **Status**: {r['Status']}")
                lines.append(f"- **Duration**: {r.get('Duration(ms)', 'N/A')}ms")

                if r.get("Payload"):
                    lines.append("- **Request Body**:")
                    try:
                        payload = json.loads(r["Payload"])
                        lines.append("  ```json")
                        lines.append(f"  {json.dumps(payload, indent=2, ensure_ascii=False)}")
                        lines.append("  ```")
                    except (json.JSONDecodeError, TypeError):
                        lines.append("  ```")
                        lines.append(f"  {r['Payload']}")
                        lines.append("  ```")

                body = r.get("Response Body", "")
                if body and body not in ("[Binary Content]", ""):
                    lines.append("- **Response Body**:")
                    try:
                        parsed_body = json.loads(body)
                        body_str = json.dumps(parsed_body, indent=2, ensure_ascii=False)
                        if len(body_str) > 500:
                            body_str = body_str[:500] + "\n  ... (truncated)"
                        lines.append("  ```json")
                        lines.append(f"  {body_str}")
                        lines.append("  ```")
                    except (json.JSONDecodeError, TypeError):
                        body_str = str(body)
                        if len(body_str) > 500:
                            body_str = body_str[:500] + "... (truncated)"
                        lines.append("  ```")
                        lines.append(f"  {body_str}")
                        lines.append("  ```")

                lines.append("")

        with open(docs_file, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))

        print(f"API docs saved to:")
        print(f"   {os.path.abspath(docs_file)}")


def compare_recordings(file1, file2, output=None):
    """
    Compare two recording files and generate a diff report.

    Args:
        file1 (str): Path to the first recording file (xlsx)
        file2 (str): Path to the second recording file (xlsx)
        output (str, optional): Output report path, defaults to printing to console
    """
    df1 = pd.read_excel(file1)
    df2 = pd.read_excel(file2)

    is_zh = "方法" in df1.columns
    method_col = "方法" if is_zh else "Method"
    url_col = "URL"
    status_col = "状态" if is_zh else "Status"
    payload_col = "载荷" if is_zh else "Payload"
    body_col = "响应体" if is_zh else "Response Body"

    def extract_endpoints(df):
        endpoints = {}
        for _, row in df.iterrows():
            try:
                parsed = urlparse(str(row[url_col]))
                key = (str(row[method_col]), parsed.netloc, parsed.path)
                if key not in endpoints:
                    endpoints[key] = row
            except Exception:
                pass
        return endpoints

    ep1 = extract_endpoints(df1)
    ep2 = extract_endpoints(df2)

    keys1 = set(ep1.keys())
    keys2 = set(ep2.keys())

    added = keys2 - keys1
    removed = keys1 - keys2
    common = keys1 & keys2

    changed = []
    for key in sorted(common):
        r1, r2 = ep1[key], ep2[key]
        diffs = []
        if str(r1[status_col]) != str(r2[status_col]):
            diffs.append(f"Status: {r1[status_col]} -> {r2[status_col]}")
        if str(r1.get(payload_col, "")) != str(r2.get(payload_col, "")):
            diffs.append("Request body changed")
        try:
            b1 = set(json.loads(str(r1[body_col])).keys())
            b2 = set(json.loads(str(r2[body_col])).keys())
            if b1 != b2:
                new_keys = b2 - b1
                del_keys = b1 - b2
                if new_keys:
                    diffs.append(f"New response fields: {', '.join(new_keys)}")
                if del_keys:
                    diffs.append(f"Removed response fields: {', '.join(del_keys)}")
        except Exception:
            pass
        if diffs:
            changed.append((key, diffs))

    lines = [
        "# API Comparison Report",
        "",
        f"File A: {os.path.basename(file1)}",
        f"File B: {os.path.basename(file2)}",
        f"Compared: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "",
        "---",
        "",
        f"## New Endpoints ({len(added)}, in B but not A)",
        ""
    ]

    if added:
        for method, domain, path in sorted(added):
            lines.append(f"- **{method}** `{domain}{path}`")
    else:
        lines.append("None")
    lines.append("")

    lines.append(f"## Removed Endpoints ({len(removed)}, in A but not B)")
    lines.append("")
    if removed:
        for method, domain, path in sorted(removed):
            lines.append(f"- **{method}** `{domain}{path}`")
    else:
        lines.append("None")
    lines.append("")

    lines.append(f"## Changed Endpoints ({len(changed)})")
    lines.append("")
    if changed:
        for (method, domain, path), diffs in changed:
            lines.append(f"### {method} `{domain}{path}`")
            for d in diffs:
                lines.append(f"- {d}")
            lines.append("")
    else:
        lines.append("None")

    report = "\n".join(lines)

    if output:
        output_dir = os.path.dirname(os.path.abspath(output))
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir)
        with open(output, "w", encoding="utf-8") as f:
            f.write(report)
        print(f"Comparison report saved to: {os.path.abspath(output)}")
    else:
        print(report)

    return report
